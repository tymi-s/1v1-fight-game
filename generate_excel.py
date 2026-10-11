"""Generator arkusza treningowego (Excel .xlsx) — wrzesień–listopad 2026.

Układ poziomy: jeden arkusz na trening (Góra A, Dół A, Góra B, Dół B, Kardio).
Wiersze = ćwiczenia i serie (stała, krótka lista), kolumny = daty treningów
(po 2 kolumny: powtórzenia i ciężar). Przewija się tylko w prawo.

Kolumny dat = dni z wpisanymi wynikami (trening_log.py) + puste, zaplanowane
treningi od PLAN_FROM do PLAN_TO. Pola do wpisywania mają format tekstowy,
żeby np. 12.5 nie zamieniało się w datę.
"""

import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from trening_log import LOG

PLAN_FROM = datetime.date(2026, 10, 11)
PLAN_TO = datetime.date(2026, 11, 30)
OUTPUT = "Trening_wrzesien_listopad_2026.xlsx"

TEXT = "@"

NAVY = "1F3A5F"
BLUE = "2E5C8A"
GREEN = "117A65"
CLUSTER = "C0392B"
WHITE = "FFFFFF"
INPUT_A = "FFFDE7"
INPUT_B = "EAF2FB"
LABEL_FILL = "F2F5F9"

thin = Side(style="thin", color="C8CED6")
thick = Side(style="medium", color="7F8C8D")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

F_TITLE = Font(name="Calibri", size=12, bold=True, color=WHITE)
F_HDR = Font(name="Calibri", size=10, bold=True, color=WHITE)
F_SUB = Font(name="Calibri", size=9, bold=True, color="1A1A1A")
F_EX = Font(name="Calibri", size=10, bold=True, color="1A1A1A")
F_CLUSTER = Font(name="Calibri", size=10, bold=True, color=CLUSTER)
F_CELL = Font(name="Calibri", size=10, color="1A1A1A")
F_TARGET = Font(name="Calibri", size=8, color="555555")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

DAY_SHORT = ["pon", "wt", "śr", "czw", "pt", "sob", "ndz"]

# Każde ćwiczenie: (nazwa, liczba_serii, cel[, "cluster"])
WORKOUTS = {
    "GA": {
        "sheet": "Góra A (ndz)",
        "title": "GÓRA A — Siła (klatka / plecy / barki / ramiona)",
        "color": NAVY,
        "exercises": [
            ("Wyciskanie leżąc (sztanga)", 4, "4–6 @ 80–85% / RPE 8"),
            ("Wiosłowanie sztangą / Pendlay row", 4, "5–6 / RPE 8"),
            ("OHP — wyciskanie żołnierskie", 3, "6–8 / RPE 8"),
            ("Podciąganie z obciążeniem", 3, "5–6 / RPE 8"),
            ("Prostowanie tricepsa (wyciąg)", 3, "8–10 / RPE 8–9"),
            ("Uginanie bicepsa (sztanga/hantle)", 3, "8–10 / RPE 8–9"),
        ],
    },
    "DA": {
        "sheet": "Dół A (wt)",
        "title": "DÓŁ A — Siła (przysiad-fokus)",
        "color": BLUE,
        "exercises": [
            ("Przysiad ze sztangą", 4, "4–6 @ 80% / RPE 8"),
            ("Martwy ciąg rumuński (RDL)", 3, "6–8 / RPE 8"),
            ("Leg press / hack squat", 3, "8–10 / RPE 8"),
            ("Uginanie nóg leżąc (dwugłowe)", 3, "10–12 / RPE 8–9"),
            ("Wspięcia na palce (łydki)", 4, "10–12 / RPE 9"),
            ("Core: Pallof press / rollout", 3, "8–10 / RPE 8"),
        ],
    },
    "GB": {
        "sheet": "Góra B (czw)",
        "title": "GÓRA B — Hipertrofia + CLUSTER podciąganie",
        "color": NAVY,
        "exercises": [
            ("Podciąganie z obciążeniem — CLUSTER", 3, "do spadku prędkości @ ~80% dodatku", "cluster"),
            ("Wyciskanie hantli na skosie", 4, "8–10 / RPE 8"),
            ("Wiosłowanie na wyciągu / chest-supported", 4, "8–12 / RPE 8–9"),
            ("Rozpiętki / cable fly", 3, "12–15 / RPE 9"),
            ("Lateral raise (barki boczne)", 4, "12–15 / RPE 9"),
            ("Uginanie młotkowe + face pull (superset)", 3, "12–15 / RPE 8–9"),
        ],
    },
    "DB": {
        "sheet": "Dół B (pt)",
        "title": "DÓŁ B — Hipertrofia (hinge-fokus)",
        "color": BLUE,
        "exercises": [
            ("Martwy ciąg", 3, "5 @ 80% / RPE 8"),
            ("Hip thrust (pośladki)", 3, "8–10 / RPE 8–9"),
            ("Wykroki / przysiad bułgarski", 3, "8–10 / noga / RPE 8"),
            ("Prostowanie nóg (czworogłowe)", 3, "12–15 / RPE 9"),
            ("Wspięcia na palce (inny kąt)", 4, "12–15 / RPE 9"),
            ("Core: hanging leg raise", 3, "10–12 / RPE 8"),
        ],
    },
}

KARDIO_PLAN = {2: "Bieg 8–10 km, Z2", 5: "Małe kardio: 8 km biegu"}


WEEK = {6: "GA", 0: None, 1: "DA", 2: "KARDIO", 3: "GB", 4: "DB", 5: "KARDIO"}


def dates_for(code):
    logged = {d for (c, d) in LOG if c == code}
    planned = set()
    d = PLAN_FROM
    while d <= PLAN_TO:
        if WEEK[d.weekday()] == code:
            planned.add(d)
        d += datetime.timedelta(days=1)
    return sorted(logged | planned)


def date_label(d):
    return f"{DAY_SHORT[d.weekday()]} {d.day:02d}.{d.month:02d}"


def style(cell, font=None, fill=None, align=None, fmt=None):
    cell.border = BORDER
    if font:
        cell.font = font
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if align:
        cell.alignment = align
    if fmt:
        cell.number_format = fmt


def fit_to_page(ws):
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def build_workout_sheet(wb, code):
    w = WORKOUTS[code]
    ws = wb.create_sheet(title=w["sheet"])
    fit_to_page(ws)
    dates = dates_for(code)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 17

    ws.merge_cells("A1:C1")
    style(ws["A1"], F_TITLE, w["color"], LEFT)
    ws["A1"] = w["title"]
    for c in ("B1", "C1"):
        style(ws[c], fill=w["color"])
    for col, text in ((1, "Ćwiczenie"), (2, "Seria"), (3, "Cel")):
        style(ws.cell(row=2, column=col, value=text), F_SUB, LABEL_FILL, CENTER)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18

    # Wiersze ćwiczeń
    row = 3
    set_rows = []  # (row, nazwa_ćwiczenia, numer_serii)
    for ex in w["exercises"]:
        name, n_sets, target = ex[0], ex[1], ex[2]
        is_cluster = len(ex) > 3
        start = row
        for s in range(1, n_sets + 1):
            label = f"K{s}" if is_cluster else s
            style(ws.cell(row=row, column=2, value=label),
                  F_CLUSTER if is_cluster else F_CELL, LABEL_FILL, CENTER)
            set_rows.append((row, name, s))
            row += 1
        end = row - 1
        if end > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
            ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
        for r in range(start, end + 1):
            style(ws.cell(row=r, column=1), fill=("FBE9E7" if is_cluster else WHITE))
            style(ws.cell(row=r, column=3), fill=WHITE)
        style(ws.cell(row=start, column=1, value=name),
              F_CLUSTER if is_cluster else F_EX, "FBE9E7" if is_cluster else WHITE, LEFT)
        style(ws.cell(row=start, column=3, value=target), F_TARGET, WHITE, CENTER)
        for r in range(start, end + 1):
            for c in (1, 2, 3):
                cell = ws.cell(row=r, column=c)
                cell.border = Border(left=thin, right=thin,
                                     top=thick if r == start else thin,
                                     bottom=thin)
    notes_row = row
    ws.merge_cells(start_row=notes_row, start_column=1, end_row=notes_row, end_column=3)
    style(ws.cell(row=notes_row, column=1, value="Notatki / samopoczucie"), F_EX, LABEL_FILL, LEFT)
    for c in (2, 3):
        style(ws.cell(row=notes_row, column=c), fill=LABEL_FILL)
    ws.row_dimensions[notes_row].height = 34

    # Kolumny dat
    for i, d in enumerate(dates):
        c_reps = 4 + i * 2
        c_kg = c_reps + 1
        fill = INPUT_A if i % 2 == 0 else INPUT_B
        ws.column_dimensions[get_column_letter(c_reps)].width = 6.5
        ws.column_dimensions[get_column_letter(c_kg)].width = 7.5

        ws.merge_cells(start_row=1, start_column=c_reps, end_row=1, end_column=c_kg)
        style(ws.cell(row=1, column=c_reps, value=date_label(d)), F_HDR, w["color"], CENTER, TEXT)
        style(ws.cell(row=1, column=c_kg), fill=w["color"])
        style(ws.cell(row=2, column=c_reps, value="Powt."), F_SUB, LABEL_FILL, CENTER)
        style(ws.cell(row=2, column=c_kg, value="kg"), F_SUB, LABEL_FILL, CENTER)

        logged = LOG.get((code, d), {})
        for r, name, s in set_rows:
            vals = logged.get(name, [])
            reps, kg = vals[s - 1] if s <= len(vals) else ("", "")
            for col, val in ((c_reps, reps), (c_kg, kg)):
                cell = ws.cell(row=r, column=col, value=val or None)
                style(cell, F_CELL, fill, CENTER, TEXT)
                if s == 1:
                    cell.border = Border(left=thin, right=thin, top=thick, bottom=thin)

        ws.merge_cells(start_row=notes_row, start_column=c_reps, end_row=notes_row, end_column=c_kg)
        style(ws.cell(row=notes_row, column=c_reps, value=logged.get("_notatka")),
              F_TARGET, fill, LEFT, TEXT)
        style(ws.cell(row=notes_row, column=c_kg), fill=fill, fmt=TEXT)

    ws.freeze_panes = "D3"
    ws.sheet_view.zoomScale = 90


def build_kardio_sheet(wb):
    ws = wb.create_sheet(title="Kardio (śr, sob)")
    fit_to_page(ws)
    dates = dates_for("KARDIO")
    labels = ["Plan", "Dystans [km]", "Czas", "Tempo [min/km]", "Notatki"]

    ws.column_dimensions["A"].width = 18
    style(ws.cell(row=1, column=1, value="KARDIO — bieg"), F_TITLE, GREEN, LEFT)
    for i, text in enumerate(labels, start=2):
        style(ws.cell(row=i, column=1, value=text), F_EX, LABEL_FILL, LEFT)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[6].height = 34

    for i, d in enumerate(dates):
        col = 2 + i
        fill = INPUT_A if i % 2 == 0 else INPUT_B
        ws.column_dimensions[get_column_letter(col)].width = 13
        style(ws.cell(row=1, column=col, value=date_label(d)), F_HDR, GREEN, CENTER, TEXT)
        style(ws.cell(row=2, column=col, value=KARDIO_PLAN[d.weekday()]), F_TARGET, LABEL_FILL, CENTER)
        for r in range(3, 7):
            style(ws.cell(row=r, column=col), F_CELL, fill, CENTER, TEXT)
    ws.freeze_panes = "B2"


def build_instrukcja(wb):
    ws = wb.create_sheet(title="Instrukcja", index=0)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    rows = [
        ("PLAN TRENINGOWY — arkusz do śledzenia progresu (wrzesień–listopad 2026)", None, True, NAVY),
        ("", None, False, None),
        ("Jak korzystać", None, True, NAVY),
        ("1.", "Każdy trening ma swoją zakładkę (Góra A, Dół A, Góra B, Dół B, Kardio).", False, None),
        ("2.", "Ćwiczenia i serie są po lewej (zablokowane). Daty treningów idą w prawo — przewijasz tylko w prawo.", False, None),
        ("3.", "Pod każdą datą dwie kolumny: Powt. i kg. Na dole wiersz na notatki.", False, None),
        ("4.", "Pola mają format tekstowy: wpisuj 12.5 albo 12,5 — nie zamieni się w datę.", False, None),
        ("5.", "Kolumny z datami mają naprzemienne kolory, żeby łatwo było trafić w właściwy trening.", False, None),
        ("6.", "Data w nagłówku to plan — jeśli trenujesz innego dnia, po prostu nadpisz ją.", False, None),
        ("", None, False, None),
        ("Tydzień", None, True, NAVY),
        ("Niedziela", "GÓRA A — Siła (klatka / plecy / barki / ramiona)", False, None),
        ("Poniedziałek", "Odpoczynek", False, None),
        ("Wtorek", "DÓŁ A — Siła (przysiad-fokus)", False, None),
        ("Środa", "Bieg 8–10 km (strefa 2)", False, None),
        ("Czwartek", "GÓRA B — Hipertrofia + CLUSTER podciąganie", False, None),
        ("Piątek", "DÓŁ B — Hipertrofia (hinge-fokus)", False, None),
        ("Sobota", "Małe kardio — 8 km biegu", False, None),
        ("", None, False, None),
        ("Cluster", "K1–K3 = klastry podciągania. W 'Powt.' wpisz ile pojedynczych powtórzeń zrobiłeś w klastrze, w 'kg' ciężar dodatku.", False, CLUSTER),
        ("Superset", "Uginanie młotkowe + face pull: w 'kg' możesz wpisać oba ciężary, np. 15 - 34.", False, None),
        ("", None, False, None),
        ("Progresja", None, True, NAVY),
        ("•", "Wszystkie serie w zakresie i RPE ≤ 8 → następnym razem +2,5 kg (sztanga) lub +1–2 powt.", False, None),
        ("•", "RPE 9–10 / ledwo domykasz zakres → zostań przy tym samym ciężarze.", False, None),
        ("•", "Klastry podciągania idą łatwo → +1–2 kg na łańcuchu.", False, None),
        ("•", "Co 6–8 tygodni deload: −40% objętości, −10% ciężaru.", False, None),
    ]
    markers = ("1.", "2.", "3.", "4.", "5.", "6.", "•")
    for i, (a, b, bold, color) in enumerate(rows, start=1):
        is_label = b is not None and a not in markers
        ca = ws.cell(row=i, column=1, value=a)
        ca.font = Font(name="Calibri", size=13 if i == 1 else 11, bold=bold or is_label,
                       color=color or "1A1A1A")
        ca.alignment = Alignment(vertical="top")
        if b:
            cb = ws.cell(row=i, column=2, value=b)
            cb.font = Font(name="Calibri", size=11, color="1A1A1A")
            cb.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.showGridLines = False


def build_progres(wb, color=NAVY):
    ws = wb.create_sheet(title="Progres (rekordy)")
    headers = ["Data", "Ćwiczenie", "Ciężar [kg]", "Powt.", "Szac. 1RM", "Notatki"]
    widths = [12, 32, 12, 8, 11, 28]
    for i, (h, wd) in enumerate(zip(headers, widths), start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
        style(ws.cell(row=1, column=i, value=h), F_HDR, color, CENTER)
    for r in range(2, 62):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            style(cell, F_CELL, INPUT_A if c != 5 else WHITE,
                  LEFT if c in (2, 6) else CENTER, TEXT if c != 5 else None)
        # NUMBERVALUE z kropką jako separatorem — działa niezależnie od ustawień regionalnych,
        # a SUBSTITUTE pozwala wpisywać zarówno 12.5 jak i 12,5.
        ws.cell(row=r, column=5).value = (
            f'=IFERROR(ROUND(_xlfn.NUMBERVALUE(SUBSTITUTE(C{r},",","."),".")'
            f'*(1+_xlfn.NUMBERVALUE(D{r},".")/30),1),"")'
        )
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    wb = Workbook()
    wb.remove(wb.active)
    build_instrukcja(wb)
    for code in ("GA", "DA", "GB", "DB"):
        build_workout_sheet(wb, code)
    build_kardio_sheet(wb)
    build_progres(wb)
    wb.save(OUTPUT)
    print(f"Zapisano: {OUTPUT}")
