"""Generator arkusza treningowego dla dziewczyny (Excel .xlsx), październik–listopad 2026.

Ten sam układ poziomy co generate_excel.py: jedna zakładka na trening,
wiersze = ćwiczenia i serie, kolumny = daty treningów (Powt. + kg), przewijanie w prawo.

Tydzień: pon Dół A · wt wolne · śr Góra A · czw Dół B · pt cardio + brzuch + biodra ·
sob Góra B · ndz wolne.
"""

import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border

from generate_excel import (
    style, fit_to_page, date_label, build_progres, get_column_letter,
    TEXT, WHITE, INPUT_A, INPUT_B, LABEL_FILL, thin, thick,
    F_TITLE, F_HDR, F_SUB, F_EX, F_CELL, F_TARGET, CENTER, LEFT,
)

PLAN_FROM = datetime.date(2026, 10, 11)
PLAN_TO = datetime.date(2026, 11, 30)
OUTPUT = "Trening_dziewczyna_pazdziernik_listopad_2026.xlsx"

PLUM = "6C2B5A"
VIOLET = "4A3B7A"
GREEN = "117A65"
SECTION_FILL = "EFE6EE"
F_FIELD = Font(name="Calibri", size=9, italic=True, color="6C2B5A")

# Blok ćwiczenia: name, target, sets [, field] — field = etykieta dodatkowego wiersza tekstowego.
# Blok "single": jeden wiersz tekstowy na datę (bez serii). Blok "section": nagłówek grupy.
DOL_A = {
    "sheet": "Dół A (pon)",
    "title": "DÓŁ A — Pośladki (akcent: hip thrust + RDL)",
    "color": PLUM,
    "blocks": [
        {"name": "Hip thrust na maszynie", "tag": "skrócenie", "target": "8–12 / RPE 8 (start ~110–120 kg)", "sets": 4},
        {"name": "RDL", "tag": "rozciągnięcie", "target": "8–10 / RPE 8", "sets": 3},
        {"name": "Przysiad bułgarski", "tag": "unilateralne", "target": "8–10 / noga, tułów w przód", "sets": 3},
        {"name": "Ławka rzymska", "tag": "izolacja", "target": "12–15 / RPE 8, plecy zaokrąglone", "sets": 3},
        {"name": "Abduktor na maszynie (pochylona w przód)", "tag": "pośl. średni — mało", "target": "12–15 / RPE 7–8, lekko", "sets": 2},
        {"name": "Łydki", "target": "12–15 / RPE 9", "sets": 3},
    ],
}

GORA_A = {
    "sheet": "Góra A (śr)",
    "title": "GÓRA A — Plecy, klatka, barki, biceps",
    "color": VIOLET,
    "blocks": [
        {"name": "Ściąganie drążka", "target": "8–12 / RPE 8", "sets": 4},
        {"name": "Wiosłowanie", "target": "10–12 / RPE 8", "sets": 3},
        {"name": "Wyciskanie na klatkę na maszynie", "target": "10–12 / RPE 8 (start ~30–40 kg)", "sets": 3},
        {"name": "Wznosy bokiem z hantlami", "target": "12–15 / RPE 9", "sets": 3},
        {"name": "Biceps — superseria lub zwykłe", "target": "superseria: hantle + 2× sztanga prosta, bez przerwy; 10–12 każda",
         "sets": 3, "field": "tryb: superseria / zwykłe"},
        {"name": "Face pull na lince", "target": "15–20 / RPE 8 (start ~15–20 kg)", "sets": 3},
    ],
}

DOL_B = {
    "sheet": "Dół B (czw)",
    "title": "DÓŁ B — Pośladki (akcent: sumo + glute bridge)",
    "color": PLUM,
    "blocks": [
        {"name": "Przysiad sumo", "tag": "rozciągnięcie", "target": "8–10 / RPE 8, szeroko, kolana na zewnątrz", "sets": 3},
        {"name": "KAS glute bridge / glute bridge", "tag": "skrócenie", "target": "10–12 / RPE 8–9, pauza 1 s u góry", "sets": 3},
        {"name": "Step up na lince", "tag": "unilateralne", "target": "10–12 / noga, nacisk na piętę", "sets": 3},
        {"name": "Kickback na lince", "tag": "izolacja", "target": "12–15 / noga, RPE 9", "sets": 3},
        {"name": "Rotacja zewn. biodra na lince / gumie", "tag": "pośl. mały — kontrola", "target": "12 / stronę, wolno, lekko", "sets": 2},
        {"name": "Łydki", "target": "12–15 / RPE 9", "sets": 3},
    ],
}

GORA_B = {
    "sheet": "Góra B (sob)",
    "title": "GÓRA B — Plecy, barki, triceps",
    "color": VIOLET,
    "blocks": [
        {"name": "Podciąganie", "target": "6–10 / RPE 8 (maszyna assist ~12 kg / guma / sama)",
         "sets": 3, "field": "wariant: guma / maszyna / sama"},
        {"name": "Wyciskanie barków", "target": "10–12 / RPE 8 (start ~12–20 kg)",
         "sets": 3, "field": "sprzęt: maszyna / hantle"},
        {"name": "Wiosłowanie", "target": "10–12 / RPE 8", "sets": 3},
        {"name": "Triceps na lince", "target": "12–15 / RPE 8–9 (start ~10–12 kg)", "sets": 3},
        {"name": "Wznosy bokiem z hantlami", "target": "12–15 / RPE 9", "sets": 3},
        {"name": "Face pull na lince", "target": "15–20 / RPE 8", "sets": 3},
    ],
}

CARDIO = {
    "sheet": "Cardio + brzuch (pt)",
    "title": "CARDIO + BRZUCH + OTWARCIE BIODER",
    "color": GREEN,
    "blocks": [
        {"section": "BIEG"},
        {"single": "Dystans [km]", "target": "5 km, tempo rozmowowe (Z2)"},
        {"single": "Czas", "target": "np. 32:10"},
        {"section": "BRZUCH"},
        {"name": "Plank", "target": "30–45 s (w 'Powt.' wpisz sekundy)", "sets": 3},
        {"name": "Dead bug", "target": "10 / stronę", "sets": 3},
        {"name": "Unoszenie kolan w zwisie / spięcia na wyciągu", "target": "12–15 / RPE 8", "sets": 3},
        {"section": "OTWARCIE BIODER (~15 min) — wpisz ✓ lub uwagi"},
        {"single": "Kontrola: hip CARs", "target": "2 × 5 powolnych krążeń / stronę"},
        {"single": "Rotacja wewn.: 90/90 lift-off", "target": "2 × 8 / stronę"},
        {"single": "Rotacja zewn.: gołąb / figure-4", "target": "2 × 30–45 s / stronę"},
        {"single": "Zgięcie: głęboki przysiad (trzymanie)", "target": "2 × 30–45 s"},
        {"single": "Bocznie: przysiad kozacki / żaba", "target": "2 × 6–8 / stronę"},
    ],
}

WEEK = {0: DOL_A, 2: GORA_A, 3: DOL_B, 4: CARDIO, 5: GORA_B}
SHEET_ORDER = [DOL_A, GORA_A, DOL_B, CARDIO, GORA_B]


def dates_for(workout):
    out = []
    d = PLAN_FROM
    while d <= PLAN_TO:
        if WEEK.get(d.weekday()) is workout:
            out.append(d)
        d += datetime.timedelta(days=1)
    return out


def date_cols(i):
    c = 4 + i * 2
    return c, c + 1


def input_pair(ws, row, i, merged, top_thick=False):
    c_reps, c_kg = date_cols(i)
    fill = INPUT_A if i % 2 == 0 else INPUT_B
    if merged:
        ws.merge_cells(start_row=row, start_column=c_reps, end_row=row, end_column=c_kg)
    for c in (c_reps, c_kg):
        cell = ws.cell(row=row, column=c)
        style(cell, F_CELL, fill, CENTER, TEXT)
        if top_thick:
            cell.border = Border(left=thin, right=thin, top=thick, bottom=thin)


def build_sheet(wb, w):
    ws = wb.create_sheet(title=w["sheet"])
    fit_to_page(ws)
    dates = dates_for(w)
    color = w["color"]

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 22

    ws.merge_cells("A1:C1")
    style(ws["A1"], F_TITLE, color, LEFT)
    ws["A1"] = w["title"]
    for c in ("B1", "C1"):
        style(ws[c], fill=color)
    for col, text in ((1, "Ćwiczenie"), (2, "Seria"), (3, "Cel")):
        style(ws.cell(row=2, column=col, value=text), F_SUB, LABEL_FILL, CENTER)
    ws.row_dimensions[1].height = 24
    for i, d in enumerate(dates):
        c_reps, c_kg = date_cols(i)
        ws.column_dimensions[get_column_letter(c_reps)].width = 6.5
        ws.column_dimensions[get_column_letter(c_kg)].width = 7.5
        ws.merge_cells(start_row=1, start_column=c_reps, end_row=1, end_column=c_kg)
        style(ws.cell(row=1, column=c_reps, value=date_label(d)), F_HDR, color, CENTER, TEXT)
        style(ws.cell(row=1, column=c_kg), fill=color)
        style(ws.cell(row=2, column=c_reps, value="Powt."), F_SUB, LABEL_FILL, CENTER)
        style(ws.cell(row=2, column=c_kg, value="kg"), F_SUB, LABEL_FILL, CENTER)

    row = 3
    for b in w["blocks"]:
        if "section" in b:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            style(ws.cell(row=row, column=1, value=b["section"]), F_SUB, SECTION_FILL, LEFT)
            for c in range(2, 4 + 2 * len(dates)):
                style(ws.cell(row=row, column=c), fill=SECTION_FILL)
            row += 1
            continue

        if "single" in b:
            style(ws.cell(row=row, column=1, value=b["single"]), F_EX, WHITE, LEFT)
            style(ws.cell(row=row, column=2, value="—"), F_CELL, LABEL_FILL, CENTER)
            style(ws.cell(row=row, column=3, value=b["target"]), F_TARGET, WHITE, CENTER)
            for i in range(len(dates)):
                input_pair(ws, row, i, merged=True)
            row += 1
            continue

        start = row
        for s in range(1, b["sets"] + 1):
            style(ws.cell(row=row, column=2, value=s), F_CELL, LABEL_FILL, CENTER)
            for i in range(len(dates)):
                input_pair(ws, row, i, merged=False, top_thick=(s == 1))
            row += 1
        end = row - 1
        name = b["name"] + (f"\n({b['tag']})" if b.get("tag") else "")
        if end > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
            ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
        for r in range(start, end + 1):
            for c in (1, 3):
                style(ws.cell(row=r, column=c), fill=WHITE)
        style(ws.cell(row=start, column=1, value=name), F_EX, WHITE, LEFT)
        style(ws.cell(row=start, column=3, value=b["target"]), F_TARGET, WHITE, CENTER)
        for c in (1, 2, 3):
            ws.cell(row=start, column=c).border = Border(left=thin, right=thin, top=thick, bottom=thin)

        if b.get("field"):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            style(ws.cell(row=row, column=1, value=f"   ↳ {b['field']}"), F_FIELD, LABEL_FILL, LEFT)
            for c in (2, 3):
                style(ws.cell(row=row, column=c), fill=LABEL_FILL)
            for i in range(len(dates)):
                input_pair(ws, row, i, merged=True)
            row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    style(ws.cell(row=row, column=1, value="Notatki / samopoczucie"), F_EX, LABEL_FILL, LEFT)
    for c in (2, 3):
        style(ws.cell(row=row, column=c), fill=LABEL_FILL)
    for i in range(len(dates)):
        input_pair(ws, row, i, merged=True)
        ws.cell(row=row, column=date_cols(i)[0]).alignment = LEFT
    ws.row_dimensions[row].height = 34

    ws.freeze_panes = "D3"
    ws.sheet_view.zoomScale = 90


def build_instrukcja(wb):
    ws = wb.create_sheet(title="Instrukcja", index=0)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 78
    head = Font(name="Calibri", size=12, bold=True, color=PLUM)
    rows = [
        ("PLAN TRENINGOWY — pośladki + plecy, redukcja (październik–listopad 2026)", None, "title"),
        ("", None, None),
        ("Jak korzystać", None, "head"),
        ("•", "Każdy trening ma swoją zakładkę. Ćwiczenia są po lewej (zablokowane), daty idą w prawo — przewijasz tylko w prawo.", None),
        ("•", "Pod każdą datą: Powt. i kg. Pola mają format tekstowy — 12.5 albo 12,5 nie zamieni się w datę.", None),
        ("•", "Szare wiersze ↳ pod ćwiczeniem: wpisz wariant (np. podciąganie: guma / maszyna / sama).", None),
        ("•", "Data w nagłówku to plan — jeśli trenujesz innego dnia, nadpisz ją.", None),
        ("", None, None),
        ("Tydzień", None, "head"),
        ("Poniedziałek", "DÓŁ A — pośladki (hip thrust, RDL, bułgarski, ławka rzymska)", None),
        ("Wtorek", "Wolne", None),
        ("Środa", "GÓRA A — ściąganie, wiosłowanie, klatka, barki, biceps, face pull", None),
        ("Czwartek", "DÓŁ B — pośladki (sumo, glute bridge, step up, kickback)", None),
        ("Piątek", "Cardio 5 km + brzuch + otwarcie bioder", None),
        ("Sobota", "GÓRA B — podciąganie, barki, wiosłowanie, triceps, face pull", None),
        ("Niedziela", "Wolne (otwarcie bioder można też zrobić tutaj albo przed dniem dołu)", None),
        ("", None, None),
        ("Dzień dołu", None, "head"),
        ("Rozciągnięcie", "RDL (pon.) / przysiad sumo (czw.)", None),
        ("Skrócenie", "Hip thrust na maszynie (pon.) / KAS glute bridge (czw.)", None),
        ("Unilateralne", "Przysiad bułgarski (pon.) / step up na lince (czw.)", None),
        ("Izolacja", "Ławka rzymska (pon.) / kickback na lince (czw.)", None),
        ("Pośl. średni/mały", "Mało i lekko: abduktor pochylona w przód (pon.), rotacja zewn. (czw.).", None),
        ("", "Pochylenie tułowia w przód na abduktorze przenosi pracę na górną część pośladka wielkiego, mniej na pośladek średni.", None),
        ("", None, None),
        ("Biceps", None, "head"),
        ("Superseria", "3 wariacje jedna po drugiej, bez przerwy = 1 seria: hantle, potem 2× sztanga prosta. W 'kg' wpisz np. 6 / 15 / 12.", None),
        ("", None, None),
        ("Otwarcie bioder", None, "head"),
        ("", "Kontrola (CARs), rotacja wewnętrzna, rotacja zewnętrzna, zgięcie, bocznie. Spokojnie, bez bólu, ~15 min. W arkuszu odhacz ✓.", None),
        ("", None, None),
        ("Progresja", None, "head"),
        ("•", "Górny zakres powtórzeń przy RPE ≤ 8 → dołóż mały ciężar (+2,5 kg maszyna / +1–2 kg hantle).", None),
        ("•", "Zakres jeszcze niepełny → zostań przy ciężarze, dokładaj powtórzenia.", None),
        ("•", "Co 6–8 tygodni lżejszy tydzień: −30–40% serii.", None),
    ]
    for i, (a, b, kind) in enumerate(rows, start=1):
        ca = ws.cell(row=i, column=1, value=a)
        if kind == "title":
            ca.font = Font(name="Calibri", size=13, bold=True, color=PLUM)
        elif kind == "head":
            ca.font = head
        else:
            ca.font = Font(name="Calibri", size=11, bold=a not in ("•", ""), color="1A1A1A")
        ca.alignment = Alignment(vertical="top")
        if b:
            cb = ws.cell(row=i, column=2, value=b)
            cb.font = Font(name="Calibri", size=11, color="1A1A1A")
            cb.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.showGridLines = False
    fit_to_page(ws)
    ws.page_setup.orientation = "portrait"


if __name__ == "__main__":
    wb = Workbook()
    wb.remove(wb.active)
    build_instrukcja(wb)
    for w in SHEET_ORDER:
        build_sheet(wb, w)
    build_progres(wb, color=PLUM)
    fit_to_page(wb["Progres (rekordy)"])
    wb.save(OUTPUT)
    print(f"Zapisano: {OUTPUT}")
